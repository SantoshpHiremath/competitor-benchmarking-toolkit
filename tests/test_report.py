import os
import tempfile

import openpyxl

from src.benchmarking import MarginBenchmark, RevenueGrowth, SegmentLeader
from src.data_sources import FinancialRecord, ProductAnnouncement
from src.reconcile import ConflictFlag, MissingDataFlag
from src.report import build_report


def _sample_inputs():
    financials = [FinancialRecord(
        company="Acme", fiscal_year=2025, reported_currency="EUR",
        revenue=100.0, ebit_margin_pct=10.0, source="Annual Report", as_of_date="2025-01-01",
    )]
    growth = [RevenueGrowth(company="Acme", fy2024_revenue=90.0, fy2025_revenue=100.0, growth_pct=11.1, rank=1)]
    margins = [MarginBenchmark(company="Acme", ebit_margin_pct=10.0, vs_peer_average_pp=0.0)]
    leaders = [SegmentLeader(segment="Motors", region="EMEA", leading_company="Acme", share_pct=30.0)]
    conflicts = [ConflictFlag(company="Acme", fiscal_year=2025, values=[("Annual Report", 100.0), ("Analyst Estimate", 110.0)], spread_pct=10.0)]
    missing = [MissingDataFlag(company="Acme", fiscal_year=2025, field="ebit_margin_pct", source="Press Release")]
    announcements = [ProductAnnouncement(company="Acme", segment="Motors", region="EMEA", announcement_date="2025-01-01", headline="Launches new line", source="Press Release")]
    return financials, growth, margins, leaders, conflicts, missing, announcements


def test_build_report_creates_real_xlsx_file():
    financials, growth, margins, leaders, conflicts, missing, announcements = _sample_inputs()
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "report.xlsx")
        build_report(path, financials, growth, margins, leaders, conflicts, missing, announcements)
        assert os.path.exists(path)
        assert os.path.getsize(path) > 0


def test_build_report_has_expected_sheets():
    financials, growth, margins, leaders, conflicts, missing, announcements = _sample_inputs()
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "report.xlsx")
        build_report(path, financials, growth, margins, leaders, conflicts, missing, announcements)
        wb = openpyxl.load_workbook(path)
        assert set(wb.sheetnames) == {
            "Summary", "Revenue Growth", "Margin Benchmark",
            "Segment Leaders", "Product Announcements", "Data Quality Flags",
        }


def test_build_report_revenue_growth_sheet_contains_data_row():
    financials, growth, margins, leaders, conflicts, missing, announcements = _sample_inputs()
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "report.xlsx")
        build_report(path, financials, growth, margins, leaders, conflicts, missing, announcements)
        wb = openpyxl.load_workbook(path)
        ws = wb["Revenue Growth"]
        assert ws["B2"].value == "Acme"
        assert ws["E2"].value == 11.1


def test_build_report_data_quality_sheet_lists_conflict_sources():
    financials, growth, margins, leaders, conflicts, missing, announcements = _sample_inputs()
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "report.xlsx")
        build_report(path, financials, growth, margins, leaders, conflicts, missing, announcements)
        wb = openpyxl.load_workbook(path)
        ws = wb["Data Quality Flags"]
        row_values = [ws.cell(row=3, column=c).value for c in range(1, 5)]
        assert row_values[0] == "Acme"
        assert "Annual Report" in row_values[3]
        assert "Analyst Estimate" in row_values[3]


def test_build_report_summary_sheet_counts_are_correct():
    financials, growth, margins, leaders, conflicts, missing, announcements = _sample_inputs()
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "report.xlsx")
        build_report(path, financials, growth, margins, leaders, conflicts, missing, announcements)
        wb = openpyxl.load_workbook(path)
        ws = wb["Summary"]
        assert ws["B3"].value == 1  # one distinct company
        assert ws["B4"].value == 1  # one conflict
        assert ws["B5"].value == 1  # one missing field
